from collections import defaultdict
from decimal import Decimal
import json
from io import BytesIO
from pathlib import Path

from django.contrib import messages
from django.contrib.auth import get_user_model
from django.contrib.auth import update_session_auth_hash
from django.contrib.auth.decorators import login_required
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import Count, Q, Sum
from django.http import FileResponse, Http404, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils.dateparse import parse_date
from django.utils import timezone
from openpyxl import load_workbook

from delegates.models import Delegate
from merchants.models import Merchant, MerchantPayment
from orders.models import Order, OrderStatus

_ORDER_STATUS_CODES = tuple(c for c, _ in OrderStatus.choices)

from .audit import log_action
from .db_backups import backup_dir, create_sqlite_backup
from .http_utils import query_first_str
from .forms import (
    AdminSetPasswordForm,
    ProfileForm,
    ProfilePasswordChangeForm,
    SiteAppearanceForm,
    UserCreateForm,
    UserEditForm,
    UserProfileRoleForm,
    UserSuperuserForm,
)
from .models import (
    AuditAction,
    AuditLog,
    Branch,
    ClientSession,
    ChatMessage,
    ChatThread,
    ExternalShipmentBatch,
    ExternalShipmentRow,
    LedgerEntry,
    OrderRequest,
    OrderRequestStatus,
    SiteAppearance,
    TreasuryAccount,
    TreasuryDirection,
    TreasuryEntry,
    TreasuryKind,
    TreasuryReason,
    UserProfile,
    AppErrorLog,
)
from .permissions import can_manage_staff_records, is_admin_level, is_super_admin, orders_queryset_for_user


def _decimal_from_post(val, default="0"):
    try:
        return Decimal(str(val or default).replace(",", ""))
    except Exception:
        return Decimal("0")


def _max_numeric_profit_from_row(values):
    best = None
    for v in values:
        if v is None or v == "":
            continue
        try:
            d = Decimal(str(v).replace(",", "").strip())
        except Exception:
            continue
        if best is None or d > best:
            best = d
    return best if best is not None else Decimal("0")


User = get_user_model()


def _create_or_get_chat_thread(user):
    thread, _ = ChatThread.objects.get_or_create(participant=user)
    return thread


def _orders_filtered(request):
    qs = Order.objects.select_related("delegate", "merchant").all()
    if not is_admin_level(request.user):
        qs = orders_queryset_for_user(request.user)
    q = query_first_str(request.GET, "q")
    if q:
        qs = qs.filter(
            Q(waybill_number__icontains=q)
            | Q(merchant__name__icontains=q)
            | Q(delegate__name__icontains=q)
        )
    status = query_first_str(request.GET, "status")
    if status in _ORDER_STATUS_CODES:
        qs = qs.filter(status=status)
    ds = query_first_str(request.GET, "date")
    day = parse_date(ds) if ds else None
    if day:
        qs = qs.filter(order_date=day)
    else:
        sds = query_first_str(request.GET, "start_date")
        eds = query_first_str(request.GET, "end_date")
        sd = parse_date(sds) if sds else None
        ed = parse_date(eds) if eds else None
        if sd:
            qs = qs.filter(order_date__gte=sd)
        if ed:
            qs = qs.filter(order_date__lte=ed)
    return qs


@login_required
def profits(request):
    if not is_admin_level(request.user):
        messages.error(request, "صفحة الأرباح للمسؤولين فقط.")
        return redirect("dashboard")

    orders = _orders_filtered(request).filter(status__in=[OrderStatus.DELIVERED, OrderStatus.ACCOUNTED])
    groups = orders.values("delegate_id", "order_date").annotate(shipping_total=Sum("shipping_price"))
    delegates = {d.pk: d for d in Delegate.objects.all()}

    total_shipping = Decimal("0")
    total_commission = Decimal("0")
    total_fixed = Decimal("0")
    total_profit = Decimal("0")
    per_delegate = defaultdict(
        lambda: {
            "shipping": Decimal("0"),
            "commission": Decimal("0"),
            "fixed": Decimal("0"),
            "profit": Decimal("0"),
        }
    )

    for g in groups:
        delegate = delegates.get(g["delegate_id"])
        shipping_total = g["shipping_total"] or Decimal("0")
        commission_rate = getattr(delegate, "commission_rate", None) or Decimal("0")
        fixed = getattr(delegate, "fixed_deduction", None) or Decimal("0")
        commission = (shipping_total * commission_rate).quantize(Decimal("0.01"))
        profit = (shipping_total - commission - fixed).quantize(Decimal("0.01"))
        total_shipping += shipping_total
        total_commission += commission
        total_fixed += fixed
        total_profit += profit
        key = delegate.name if delegate else "غير محدد"
        per_delegate[key]["shipping"] += shipping_total
        per_delegate[key]["commission"] += commission
        per_delegate[key]["fixed"] += fixed
        per_delegate[key]["profit"] += profit

    internal_treasury = TreasuryEntry.objects.filter(account__kind=TreasuryKind.INTERNAL)
    external_treasury = TreasuryEntry.objects.filter(account__kind=TreasuryKind.EXTERNAL)
    ext_ship_profit = ExternalShipmentRow.objects.aggregate(s=Sum("profit_amount"))["s"] or Decimal("0")

    totals = {
        "orders_count": orders.count(),
        "shipping_total": total_shipping,
        "commission_total": total_commission,
        "fixed_total": total_fixed,
        "profit_total": total_profit,
        "treasury_internal_in": internal_treasury.filter(direction=TreasuryDirection.IN).aggregate(
            s=Sum("amount")
        )["s"]
        or Decimal("0"),
        "treasury_external_in": external_treasury.filter(direction=TreasuryDirection.IN).aggregate(
            s=Sum("amount")
        )["s"]
        or Decimal("0"),
        "external_shipments_profit": ext_ship_profit,
    }

    per_delegate_rows = [
        {"delegate_name": name, **vals}
        for name, vals in sorted(per_delegate.items(), key=lambda x: x[0])
    ]
    profit_chart_labels = [r["delegate_name"] for r in per_delegate_rows]
    profit_chart_values = [float(r["profit"]) for r in per_delegate_rows]

    return render(
        request,
        "core/profits.html",
        {
            "totals": totals,
            "per_delegate_rows": per_delegate_rows,
            "profit_chart_labels_json": json.dumps(profit_chart_labels, ensure_ascii=False),
            "profit_chart_values_json": json.dumps(profit_chart_values),
        },
    )

@login_required
def treasury(request):
    # ================= POST ACTIONS =================
    if request.method == "POST":
        action = request.POST.get("action")

        if action == "entry":
            if request.POST.get("confirm_entry") != "1":
                messages.error(request, "يرجى تأكيد الحركة (المربع والزر) قبل التسجيل.")
                return redirect("core:treasury")

            acc = get_object_or_404(
                TreasuryAccount,
                pk=request.POST.get("account_id")
            )
            amount = Decimal(str(request.POST.get("amount") or "0"))

            raw_dir = (request.POST.get("direction") or "").strip()
            direction = raw_dir if raw_dir in TreasuryDirection.values else TreasuryDirection.IN

            raw_reason = (request.POST.get("reason") or "").strip()
            if raw_reason not in TreasuryReason.values:
                raw_reason = TreasuryReason.MANUAL

            if amount > 0:
                TreasuryEntry.objects.create(
                    account=acc,
                    amount=amount,
                    direction=direction,
                    reason=raw_reason,
                    notes=(request.POST.get("notes") or "").strip(),
                    created_by=request.user,
                )
                messages.success(request, "تم إضافة حركة خزنة.")
            else:
                messages.error(request, "أدخل مبلغًا أكبر من صفر.")

            return redirect("core:treasury")

        if action == "close":
            if request.POST.get("confirm_close") != "1":
                messages.error(request, "يرجى تأكيد تصفية الخزنة قبل المتابعة.")
                return redirect("core:treasury")
            total_in = TreasuryEntry.objects.filter(
                direction=TreasuryDirection.IN
            ).aggregate(s=Sum("amount"))["s"] or Decimal("0")

            total_out = TreasuryEntry.objects.filter(
                direction=TreasuryDirection.OUT
            ).aggregate(s=Sum("amount"))["s"] or Decimal("0")

            messages.success(
                request,
                f"تصفية حالية: وارد {total_in} - منصرف {total_out} = صافي {total_in - total_out}"
            )
            return redirect("core:treasury")

    # ================= PERMISSION =================
    if not is_admin_level(request.user):
        messages.error(request, "صفحة الخزنة للمسؤولين فقط.")
        return redirect("dashboard")

    # ================= FILTERS =================
    dts = query_first_str(request.GET, "date")
    sts = query_first_str(request.GET, "start_date")
    ens = query_first_str(request.GET, "end_date")

    parsed_date = parse_date(dts) if dts else None
    parsed_start = parse_date(sts) if sts else None
    parsed_end = parse_date(ens) if ens else None

    # ================= BASE QUERY =================
    base_entries = TreasuryEntry.objects.select_related("account", "order").all()

    if parsed_date:
        base_entries = base_entries.filter(entry_date=parsed_date)
    else:
        if parsed_start:
            base_entries = base_entries.filter(entry_date__gte=parsed_start)
        if parsed_end:
            base_entries = base_entries.filter(entry_date__lte=parsed_end)

    # ================= TABLE =================
    entries = base_entries.order_by("-entry_date", "-id")[:200]

    # ================= CHART =================
    chart_rows = (
        base_entries.values("entry_date")
        .annotate(
            in_total=Sum("amount", filter=Q(direction=TreasuryDirection.IN)),
            out_total=Sum("amount", filter=Q(direction=TreasuryDirection.OUT)),
        )
        .order_by("entry_date")
    )

    chart_labels = [str(r["entry_date"]) for r in chart_rows]
    chart_in = [float(r["in_total"] or 0) for r in chart_rows]
    chart_out = [float(r["out_total"] or 0) for r in chart_rows]

    # ================= PAYMENTS =================
    payments = MerchantPayment.objects.all()

    if parsed_date:
        payments = payments.filter(payment_date=parsed_date)
    else:
        if parsed_start:
            payments = payments.filter(payment_date__gte=parsed_start)
        if parsed_end:
            payments = payments.filter(payment_date__lte=parsed_end)

    payment_summary = payments.aggregate(
        in_total=Sum("amount", filter=Q(direction="in")),
        out_total=Sum("amount", filter=Q(direction="out")),
    )

    # ================= ACCOUNTS =================
    accounts = TreasuryAccount.objects.filter(is_active=True)

    # ================= RENDER =================
    return render(
        request,
        "core/treasury.html",
        {
            "entries": entries,
            "accounts": accounts,
            "payment_summary": payment_summary,
            "treasury_reasons": TreasuryReason.choices,
            "chart_labels_json": json.dumps(chart_labels, ensure_ascii=False),
            "chart_in_json": json.dumps(chart_in),
            "chart_out_json": json.dumps(chart_out),
        },
    )


@login_required
def activity(request):
    qs = _orders_filtered(request)
    td_rows = (
        qs.exclude(delegate__isnull=True)
        .values("delegate__name")
        .annotate(c=Count("id"))
        .order_by("-c")[:8]
    )
    tm_rows = (
        qs.exclude(merchant__isnull=True)
        .values("merchant__name")
        .annotate(c=Count("id"))
        .order_by("-c")[:8]
    )
    top_delegates = [
        {"name": (row.get("delegate__name") or "").strip() or "—", "c": row["c"]}
        for row in td_rows
    ]
    top_merchants = [
        {"name": (row.get("merchant__name") or "").strip() or "—", "c": row["c"]}
        for row in tm_rows
    ]

    dl = [r["name"] for r in top_delegates]
    dv = [int(r["c"]) for r in top_delegates]
    ml = [r["name"] for r in top_merchants]
    mv = [int(r["c"]) for r in top_merchants]

    return render(
        request,
        "core/activity.html",
        {
            "top_delegates": top_delegates,
            "top_merchants": top_merchants,
            "delegate_chart_labels_json": json.dumps(dl, ensure_ascii=False),
            "delegate_chart_values_json": json.dumps(dv),
            "merchant_chart_labels_json": json.dumps(ml, ensure_ascii=False),
            "merchant_chart_values_json": json.dumps(mv),
        },
    )


@login_required
def devices_list(request):
    devices = ClientSession.objects.select_related("user").all()
    return render(request, "core/devices.html", {"devices": devices})


@login_required
def settings_site(request):
    if not is_admin_level(request.user):
        messages.error(request, "غير مسموح بتعديل إعدادات المظهر.")
        return redirect("dashboard")

    obj, _ = SiteAppearance.objects.get_or_create(pk=1)

    if request.method == "POST":
        form = SiteAppearanceForm(request.POST, request.FILES, instance=obj)
        if form.is_valid():
            form.save()
            messages.success(request, "تم حفظ الإعدادات")
            return redirect("core:settings")
    else:
        form = SiteAppearanceForm(instance=obj)

    return render(request, "core/settings_site.html", {"form": form})


@login_required
def external_shipments(request):
    if not is_admin_level(request.user):
        messages.error(request, "غير مسموح.")
        return redirect("dashboard")

    if request.method == "POST" and request.FILES.get("file"):
        upload = request.FILES["file"]
        buf = BytesIO(upload.read())
        try:
            wb = load_workbook(buf, data_only=True)
        except Exception:
            messages.error(request, "ملف Excel غير صالح.")
            return redirect("core:external")

        sheet = wb.active
        if sheet is None and wb.sheetnames:
            sheet = wb[wb.sheetnames[0]]
        if sheet is None:
            messages.error(request, "ملف Excel بدون أورقة عمل.")
            return redirect("core:external")

        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        headers = [str(h).strip() if h is not None else "" for h in header_row]

        batch = ExternalShipmentBatch.objects.create(
            imported_by=request.user,
            source_name=getattr(upload, "name", "") or "",
        )
        row_count = 0
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if row is None or all(v is None or str(v).strip() == "" for v in row):
                continue
            cells = list(row)
            pairs = {}
            for i, h in enumerate(headers):
                if not h:
                    continue
                val = cells[i] if i < len(cells) else None
                pairs[h] = val if val is not None else ""
            profit = _max_numeric_profit_from_row(cells)
            ExternalShipmentRow.objects.create(batch=batch, row_data=pairs, profit_amount=profit)
            row_count += 1

        messages.success(request, f"تم الاستيراد: دفعة #{batch.pk} ({row_count} صف).")
        return redirect("core:external")

    batches = ExternalShipmentBatch.objects.all().order_by("-imported_at")
    return render(request, "core/external.html", {"batches": batches})


@login_required
def order_requests_list(request):
    if not is_admin_level(request.user):
        messages.error(request, "صفحة اعتماد الطلبات للمسؤولين فقط.")
        return redirect("dashboard")

    delegates = Delegate.objects.filter(is_active=True).order_by("name")

    if request.method == "POST":
        action = (request.POST.get("action") or "").strip()
        req_id = request.POST.get("request_id")
        try:
            rid = int(req_id)
        except (TypeError, ValueError):
            rid = 0
        order_req = get_object_or_404(OrderRequest.objects.select_related("merchant"), pk=rid)

        if order_req.status != OrderRequestStatus.PENDING:
            messages.error(request, "هذا الطلب تمت معالجته مسبقًا.")
            return redirect("core:order_requests")

        if action == "reject":
            order_req.status = OrderRequestStatus.REJECTED
            order_req.save(update_fields=["status", "updated_at"])
            log_action(request, AuditAction.UPDATE, order_req, "رفض طلب شحنة")
            messages.success(request, "تم رفض الطلب.")
            return redirect("core:order_requests")

        if action == "approve":
            try:
                did = int(request.POST.get("delegate_id") or 0)
            except (TypeError, ValueError):
                did = 0
            delegate = get_object_or_404(Delegate, pk=did, is_active=True)
            with transaction.atomic():
                order = Order(
                    merchant=order_req.merchant,
                    delegate=delegate,
                    customer_name=order_req.customer_name,
                    customer_phone=order_req.customer_phone,
                    customer_address=order_req.customer_address,
                    product_price=order_req.product_price,
                    shipping_price=order_req.shipping_price,
                    notes=order_req.notes,
                    status=OrderStatus.IN_TRANSIT,
                    created_by=request.user,
                    updated_by=request.user,
                    owned_by=request.user,
                )
                setattr(order, "_save_user", request.user)
                order.save()
                order_req.created_order = order
                order_req.assigned_delegate = delegate
                order_req.status = OrderRequestStatus.APPROVED
                order_req.save(
                    update_fields=[
                        "created_order",
                        "assigned_delegate",
                        "status",
                        "updated_at",
                    ]
                )
            log_action(request, AuditAction.UPDATE, order_req, "اعتماد طلب شحنة وإنشاء أوردر")
            messages.success(request, "تم اعتماد الطلب وإنشاء الشحنة.")
            return redirect("core:order_requests")

        messages.error(request, "إجراء غير معروف.")
        return redirect("core:order_requests")

    reqs = (
        OrderRequest.objects.select_related("merchant", "requester", "assigned_delegate", "created_order")
        .all()
        .order_by("-created_at")
    )
    return render(
        request,
        "core/order_requests.html",
        {"requests": reqs, "delegates": delegates},
    )


@login_required
def order_request_new(request):
    if request.method == "POST":
        customer_name = (request.POST.get("customer_name") or "").strip()
        if not customer_name:
            messages.error(request, "اسم العميل مطلوب.")
        else:
            OrderRequest.objects.create(
                requester=request.user,
                customer_name=customer_name,
                customer_phone=(request.POST.get("customer_phone") or "").strip(),
                customer_address=(request.POST.get("customer_address") or "").strip(),
                product_price=_decimal_from_post(request.POST.get("product_price")),
                shipping_price=_decimal_from_post(request.POST.get("shipping_price")),
                notes=(request.POST.get("notes") or "").strip(),
                status=OrderRequestStatus.PENDING,
            )
            messages.success(request, "تم إرسال الطلب")
            return redirect("core:order_request_new")

    my_requests = OrderRequest.objects.filter(requester=request.user).order_by("-created_at")[:20]
    return render(request, "core/order_request_new.html", {"my_requests": my_requests})


@login_required
def chat_center(request):
    user = request.user
    admin = is_admin_level(user)

    if request.method == "POST":
        body = (request.POST.get("body") or "").strip()
        tid = request.POST.get("thread_id")
        try:
            tid_int = int(tid)
        except (TypeError, ValueError):
            tid_int = 0
        thread = get_object_or_404(ChatThread.objects.select_related("participant"), pk=tid_int)
        if admin or thread.participant_id == user.id:
            if body:
                ChatMessage.objects.create(thread=thread, sender=user, body=body)
                thread.save(update_fields=["updated_at"])
            return redirect(f"{reverse('core:chat')}?thread={thread.pk}")
        messages.error(request, "غير مسموح بإرسال رسائل في هذه المحادثة.")
        return redirect("core:chat")

    if admin:
        threads = ChatThread.objects.select_related("participant").order_by("-updated_at")
    else:
        own = _create_or_get_chat_thread(user)
        threads = ChatThread.objects.filter(pk=own.pk).select_related("participant")

    selected_thread = None
    qs = query_first_str(request.GET, "thread")
    if qs:
        try:
            cand = get_object_or_404(ChatThread.objects.select_related("participant"), pk=int(qs))
        except ValueError:
            cand = None
        else:
            if admin or cand.participant_id == user.id:
                selected_thread = cand
    if selected_thread is None and threads.count() == 1:
        selected_thread = threads.first()

    chat_messages = []
    if selected_thread is not None:
        chat_messages = list(
            ChatMessage.objects.filter(thread=selected_thread)
            .select_related("sender")
            .order_by("created_at", "id")
        )

    return render(
        request,
        "core/chat_center.html",
        {
            "threads": threads,
            "selected_thread": selected_thread,
            "chat_messages": chat_messages,
        },
    )


@login_required
def audit_list(request):
    if not is_admin_level(request.user):
        messages.error(request, "لا صلاحية لعرض سجل العمليات.")
        return redirect("dashboard")
    logs = AuditLog.objects.select_related("user").all().order_by("-created_at")
    return render(request, "core/audit.html", {"logs": logs})


@login_required
def users_list(request):
    if not can_manage_staff_records(request.user):
        messages.error(request, "لا صلاحية لعرض المستخدمين.")
        return redirect("dashboard")
    users = []
    for u in User.objects.all().select_related("profile").order_by("username"):
        role_label = "—"
        try:
            role_label = u.profile.get_role_display()
        except UserProfile.DoesNotExist:
            pass
        users.append({"user": u, "role_label": role_label})
    return render(request, "core/users_list.html", {"users": users})


@login_required
def user_add(request):
    if not is_admin_level(request.user):
        messages.error(request, "إضافة مستخدمين متاحة للمسؤولين فقط.")
        return redirect("core:users")
    if request.method == "POST":
        form = UserCreateForm(request.POST)
        if form.is_valid():
            new_user = form.save()
            prof = new_user.profile
            prof.role = form.cleaned_data["role"]
            prof.save(update_fields=["role"])
            messages.success(request, "تم إضافة المستخدم")
            return redirect("core:users")
    else:
        form = UserCreateForm()

    return render(
        request,
        "core/user_form.html",
        {"form": form, "mode": "add"},
    )


@login_required
def user_edit(request, user_id):
    if not is_admin_level(request.user):
        messages.error(request, "تعديل المستخدمين متاح للمسؤولين فقط.")
        return redirect("core:users")
    target_user = get_object_or_404(User, pk=user_id)
    profile, _ = UserProfile.objects.get_or_create(user=target_user)

    linked_merchant = Merchant.objects.filter(linked_user=target_user).first()
    linked_delegate = Delegate.objects.filter(linked_user=target_user).first()

    form = UserEditForm(instance=target_user)
    profile_form = UserProfileRoleForm(instance=profile)
    password_form = AdminSetPasswordForm(target_user)
    super_form = UserSuperuserForm(instance=target_user) if request.user.is_superuser else None

    if request.method == "POST":
        if request.POST.get("password_form"):
            password_form = AdminSetPasswordForm(target_user, request.POST)
            if password_form.is_valid():
                password_form.save()
                messages.success(request, "تم تحديث كلمة المرور.")
                return redirect("core:user_edit", user_id=target_user.pk)
        elif request.POST.get("edit_form"):
            form = UserEditForm(request.POST, instance=target_user)
            profile_form = UserProfileRoleForm(request.POST, instance=profile)
            if request.user.is_superuser:
                super_form = UserSuperuserForm(request.POST, instance=target_user)
            ok = form.is_valid() and profile_form.is_valid()
            if super_form is not None:
                ok = ok and super_form.is_valid()
            if ok:
                form.save()
                profile_form.save()
                if super_form is not None:
                    super_form.save()
                messages.success(request, "تم حفظ التعديلات.")
                return redirect("core:users")

    return render(
        request,
        "core/user_form.html",
        {
            "form": form,
            "profile_form": profile_form,
            "password_form": password_form,
            "super_form": super_form,
            "mode": "edit",
            "target_user": target_user,
            "linked_merchant": linked_merchant,
            "linked_delegate": linked_delegate,
        },
    )


@login_required
def user_soft_delete(request, user_id):
    if not can_manage_staff_records(request.user):
        messages.error(request, "لا صلاحية لتعطيل المستخدمين.")
        return redirect("dashboard")
    target_user = get_object_or_404(User, pk=user_id)
    if target_user.pk == request.user.pk:
        messages.error(request, "لا يمكنك تعطيل حسابك الحالي.")
        return redirect("core:users")
    if target_user.is_superuser and not request.user.is_superuser:
        messages.error(request, "يمكن لمدير النظام فقط تعديل حسابات المدير العام.")
        return redirect("core:users")
    if request.method != "POST":
        return redirect("core:users")
    target_user.is_active = False
    target_user.save(update_fields=["is_active"])
    log_action(
        request,
        AuditAction.DELETE,
        target_user,
        f"تعطيل مستخدم (حذف منطقي): {target_user.username}",
    )
    messages.success(request, "تم تعطيل الحساب.")
    return redirect("core:users")


@login_required
def ledger_summary(request):
    """موازنة بصرية بين الخزنة القائمة ودفتر الحركة الجديد؛ للمدير العام فقط."""

    if not is_super_admin(request.user):
        messages.error(request, "هذه الصفحة للمدير العام فقط.")
        return redirect("dashboard")

    treasury_internal_in = (
        TreasuryEntry.objects.filter(
            direction=TreasuryDirection.IN,
            account__kind=TreasuryKind.INTERNAL,
        ).aggregate(s=Sum("amount"))["s"]
        or Decimal("0")
    )

    ledger_qs = LedgerEntry.objects.all()
    ledger_sum = ledger_qs.aggregate(s=Sum("amount"))["s"] or Decimal("0")
    ledger_count = ledger_qs.count()
    ledger_by_type = (
        ledger_qs.values("entry_type").annotate(total=Sum("amount"), count=Count("id")).order_by("entry_type")
    )

    return render(
        request,
        "core/ledger_summary.html",
        {
            "treasury_internal_in": treasury_internal_in,
            "ledger_sum": ledger_sum,
            "ledger_count": ledger_count,
            "ledger_by_type": ledger_by_type,
            "branch_count": Branch.objects.count(),
        },
    )


@login_required
def profile(request):
    password_form = ProfilePasswordChangeForm(user=request.user)
    if request.method == "POST":
        if request.POST.get("password_form"):
            form = ProfileForm(instance=request.user)
            password_form = ProfilePasswordChangeForm(user=request.user, data=request.POST)
            if password_form.is_valid():
                user = password_form.save()
                update_session_auth_hash(request, user)
                messages.success(request, "تم تحديث كلمة المرور.")
                return redirect("core:profile")
        else:
            form = ProfileForm(request.POST, instance=request.user)
            if form.is_valid():
                form.save()
                messages.success(request, "تم تحديث البيانات")
                return redirect("core:profile")
    else:
        form = ProfileForm(instance=request.user)

    return render(request, "core/profile.html", {"form": form, "password_form": password_form})


@login_required
def superadmin_orders_control(request):
    if not is_super_admin(request.user):
        messages.error(request, "هذه الصفحة للمدير العام فقط.")
        return redirect("dashboard")

    if request.method == "POST":
        action = (request.POST.get("action") or "").strip()

        if action == "create_backup":
            backup_path = create_sqlite_backup("superadmin")
            messages.success(request, f"تم إنشاء نسخة احتياطية: {backup_path.name}")
            return redirect("core:superadmin_orders_control")

        if action == "restore_order":
            try:
                oid = int(request.POST.get("order_id") or 0)
            except (TypeError, ValueError):
                oid = 0
            order = get_object_or_404(Order.all_objects.select_related("merchant", "delegate"), pk=oid)
            if order.deleted_at is None:
                messages.info(request, "الأوردر ظاهر بالفعل.")
                return redirect(request.get_full_path())
            order.deleted_at = None
            order.deleted_by = None
            order.updated_by = request.user
            setattr(order, "_save_user", request.user)
            order.save(update_fields=["deleted_at", "deleted_by", "updated_by", "updated_at"])
            log_action(request, AuditAction.UPDATE, order, f"استرجاع أوردر محذوف: {order.waybill_number}")
            messages.success(request, "تم استرجاع الأوردر بنجاح.")
            return redirect(request.get_full_path())

        if action == "restore_all_deleted":
            n = Order.all_objects.filter(deleted_at__isnull=False).update(
                deleted_at=None,
                deleted_by=None,
                updated_by=request.user,
            )
            log_action(request, AuditAction.UPDATE, None, f"استرجاع جماعي للأوردرات المحذوفة: {n}")
            messages.success(request, f"تم استرجاع {n} أوردر.")
            return redirect(request.get_full_path())

    q = query_first_str(request.GET, "q")
    show = (request.GET.get("show") or "all").strip()

    orders = Order.all_objects.select_related("merchant", "delegate", "deleted_by", "created_by")
    if q:
        orders = orders.filter(
            Q(waybill_number__icontains=q)
            | Q(customer_name__icontains=q)
            | Q(customer_phone__icontains=q)
            | Q(merchant__name__icontains=q)
            | Q(delegate__name__icontains=q)
        )

    if show == "deleted":
        orders = orders.filter(deleted_at__isnull=False)
    elif show == "orphaned":
        orders = orders.filter(Q(merchant__isnull=True) | Q(delegate__isnull=True))

    backups = sorted(backup_dir().glob("*.bak"), key=lambda p: p.stat().st_mtime, reverse=True)[:30]
    backup_rows = [
        {
            "name": p.name,
            "size": p.stat().st_size,
            "modified_at": timezone.datetime.fromtimestamp(p.stat().st_mtime),
        }
        for p in backups
    ]

    page_obj = Paginator(orders.order_by("-order_date", "-id"), 50).get_page(query_first_str(request.GET, "page") or 1)

    return render(
        request,
        "core/superadmin_orders_control.html",
        {
            "page_obj": page_obj,
            "show": show,
            "backup_rows": backup_rows,
            "deleted_count": Order.all_objects.filter(deleted_at__isnull=False).count(),
            "orphaned_count": Order.all_objects.filter(Q(merchant__isnull=True) | Q(delegate__isnull=True)).count(),
            "all_orders_count": Order.all_objects.count(),
            "visible_orders_count": Order.objects.count(),
        },
    )


@login_required
def error_logs_center(request):
    if not is_super_admin(request.user):
        messages.error(request, "هذه الصفحة للمدير العام فقط.")
        return redirect("dashboard")

    if request.method == "POST" and request.POST.get("action") == "clear":
        n, _ = AppErrorLog.objects.all().delete()
        messages.success(request, f"تم مسح {n} سجل خطأ.")
        return redirect("core:error_logs_center")

    q = query_first_str(request.GET, "q")
    entries = _build_terminal_entries(search=q, limit=250)
    latest_error_id = AppErrorLog.objects.order_by("-id").values_list("id", flat=True).first() or 0
    return render(
        request,
        "core/error_logs_center.html",
        {
            "entries": entries,
            "latest_error_id": latest_error_id,
        },
    )


def _build_terminal_entries(search: str = "", limit: int = 250):
    err_qs = AppErrorLog.objects.select_related("user").all()
    audit_qs = AuditLog.objects.select_related("user").all()
    if search:
        err_qs = err_qs.filter(Q(error_type__icontains=search) | Q(message__icontains=search) | Q(path__icontains=search))
        audit_qs = audit_qs.filter(Q(message__icontains=search) | Q(path__icontains=search) | Q(model_name__icontains=search))

    error_rows = []
    for e in err_qs.order_by("-created_at", "-id")[:limit]:
        user_label = e.user.username if e.user else "anon"
        msg = f"[{e.created_at:%Y-%m-%d %H:%M:%S}] ERROR {e.error_type or 'Exception'} user={user_label} path={e.path or '-'} :: {e.message or '-'}"
        error_rows.append(
            {
                "created_at": e.created_at,
                "level": "error",
                "line": msg,
                "error_id": e.id,
            }
        )

    info_rows = []
    for a in audit_qs.order_by("-created_at", "-id")[:limit]:
        user_label = a.user.username if a.user else "anon"
        level = "warn" if (a.extra or {}).get("status_code", 200) and int((a.extra or {}).get("status_code", 200)) >= 400 else "info"
        msg = f"[{a.created_at:%Y-%m-%d %H:%M:%S}] {level.upper()} user={user_label} path={a.path or '-'} :: {a.message or '-'}"
        info_rows.append(
            {
                "created_at": a.created_at,
                "level": level,
                "line": msg,
                "error_id": 0,
            }
        )

    merged = sorted(error_rows + info_rows, key=lambda r: r["created_at"], reverse=True)[:limit]
    return merged


@login_required
def error_logs_feed(request):
    if not is_super_admin(request.user):
        return JsonResponse({"ok": False, "error": "forbidden"}, status=403)

    q = query_first_str(request.GET, "q")
    try:
        last_error_id = int(request.GET.get("last_error_id") or 0)
    except (TypeError, ValueError):
        last_error_id = 0

    entries = _build_terminal_entries(search=q, limit=120)
    latest_error_id = AppErrorLog.objects.order_by("-id").values_list("id", flat=True).first() or 0
    new_errors_count = AppErrorLog.objects.filter(id__gt=last_error_id).count() if last_error_id > 0 else 0
    payload = [
        {
            "level": e["level"],
            "line": e["line"],
            "error_id": e["error_id"],
        }
        for e in entries
    ]
    return JsonResponse(
        {
            "ok": True,
            "entries": payload,
            "latest_error_id": latest_error_id,
            "new_errors_count": new_errors_count,
        }
    )


@login_required
def db_backup_download(request, filename: str):
    if not is_super_admin(request.user):
        raise Http404
    file_path = backup_dir() / Path(filename).name
    if not file_path.exists() or not file_path.is_file():
        raise Http404
    return FileResponse(open(file_path, "rb"), as_attachment=True, filename=file_path.name)