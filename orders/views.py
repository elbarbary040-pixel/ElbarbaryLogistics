import json
import re
from datetime import timedelta
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from uuid import uuid4
from urllib.parse import quote

from django.conf import settings
from django.contrib import messages
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required
from django.core.paginator import Paginator
from django.db.models import Count, Max, Min, Q, Sum
from django.http import HttpRequest, HttpResponse, HttpResponseRedirect
from django.shortcuts import get_object_or_404, redirect, render
from django.http import JsonResponse
from django.urls import reverse
from django.utils import timezone
from django.utils.dateparse import parse_date
from django.utils.http import url_has_allowed_host_and_scheme
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
import qrcode

from core.audit import log_action
from core.db_backups import create_sqlite_backup
from core.http_utils import query_first_str, query_first_value
from core.models import AuditAction
from core.permissions import can_manage_staff_records, is_admin_level, is_super_admin, orders_queryset_for_user
from delegates.models import Delegate
from merchants.models import Merchant

from .forms import OrderForm, OrderImportMappingForm, OrderImportUploadForm
from .models import Order, OrderStatus
from .services.settlement import compute_order_settlement
from .whatsapp_ar import build_bulk_whatsapp_text_ar
from .qr_tokens import sign_order_pk, unsign_order_token

_ORDER_STATUS_CODES = tuple(c for c, _ in OrderStatus.choices)
_ORDER_STATUS_LABELS = dict(OrderStatus.choices)
_DASH_LEGACY_THREE = frozenset(
    {
        OrderStatus.DELIVERED,
        OrderStatus.IN_TRANSIT,
        OrderStatus.ACCOUNTED,
    }
)

_CHART_HEX = {
    OrderStatus.DELIVERED: "#198754",
    OrderStatus.IN_TRANSIT: "#ffc107",
    OrderStatus.ACCOUNTED: "#dc3545",
    OrderStatus.PARTIALLY_DELIVERED: "#0dcaf0",
    OrderStatus.PARTIALLY_DELIVERED_WITH_RETURN: "#fd7e14",
    OrderStatus.RETURNED: "#6f42c1",
    OrderStatus.POSTPONED: "#adb5bd",
}

STUB_PK = 987654321

WAYBILL_RE = re.compile(r"^BRB\d+$")


def _to_str(v):
    return str(v).strip() if v is not None else ""


def _to_decimal(v):
    try:
        return Decimal(str(v).replace(",", "")) if v else Decimal("0")
    except Exception:
        return Decimal("0")


def _to_date(v):
    if hasattr(v, "date"):
        return v.date()
    return parse_date(str(v)) if v else timezone.now().date()


def _get(row, idx):
    return row[idx] if idx is not None and idx < len(row) else None


def _merchant_autocomplete_qs(request: HttpRequest):
    q = query_first_str(request.GET, "q")
    qs = Merchant.objects.all()
    if q:
        qs = qs.filter(Q(name__icontains=q) | Q(phone__icontains=q))
    return qs.order_by("name")[:50]


def _delegate_autocomplete_qs(request: HttpRequest):
    q = query_first_str(request.GET, "q")
    qs = Delegate.objects.all()
    if q:
        qs = qs.filter(Q(name__icontains=q) | Q(phone__icontains=q))
    return qs.order_by("name")[:50]


@login_required
def merchant_autocomplete(request: HttpRequest):
    if not can_manage_staff_records(request.user):
        return JsonResponse({"results": []}, status=403)
    rows = _merchant_autocomplete_qs(request)
    results = [
        {"id": m.pk, "text": m.name + (f" — {m.phone}" if m.phone else "")}
        for m in rows
    ]
    return JsonResponse({"results": results})


@login_required
def delegate_autocomplete(request: HttpRequest):
    if not can_manage_staff_records(request.user):
        return JsonResponse({"results": []}, status=403)
    rows = _delegate_autocomplete_qs(request)
    results = [
        {"id": d.pk, "text": d.name + (f" — {d.phone}" if d.phone else "")}
        for d in rows
    ]
    return JsonResponse({"results": results})


def _apply_order_filters(qs, request: HttpRequest):
    q = query_first_str(request.GET, "q")
    if q:
        qs = qs.filter(
            Q(waybill_number__icontains=q)
            | Q(external_waybill__icontains=q)
            | Q(customer_name__icontains=q)
            | Q(customer_phone__icontains=q)
            | Q(merchant__name__icontains=q)
            | Q(merchant__phone__icontains=q)
            | Q(delegate__name__icontains=q)
            | Q(delegate__phone__icontains=q)
        )
    status = query_first_str(request.GET, "status")
    if status in _ORDER_STATUS_CODES:
        qs = qs.filter(status=status)
    ds = query_first_str(request.GET, "date")
    day = parse_date(ds) if ds else None
    if day:
        qs = qs.filter(order_date=day)
    else:
        sds = query_first_str(request.GET, "start_date")
        eds = query_first_str(request.GET, "end_date")
        sd = parse_date(sds) if sds else None
        ed = parse_date(eds) if eds else None
        if sd:
            qs = qs.filter(order_date__gte=sd)
        if ed:
            qs = qs.filter(order_date__lte=ed)
    return qs


def login_view(request: HttpRequest) -> HttpResponse:
    if request.user.is_authenticated:
        return redirect("dashboard")

    if request.method == "POST":
        username = query_first_str(request.POST, "username")
        password = query_first_value(request.POST, "password")
        next_url = query_first_str(request.POST, "next")
        user = authenticate(request, username=username, password=password)
        if user is not None:
            login(request, user)
            if next_url and url_has_allowed_host_and_scheme(
                next_url,
                allowed_hosts={request.get_host()},
                require_https=request.is_secure(),
            ):
                return redirect(next_url)
            return redirect("dashboard")
        messages.error(request, "بيانات الدخول غير صحيحة.")

    return render(
        request,
        "login.html",
        {"next": query_first_str(request.GET, "next")},
    )


@login_required
def logout_view(request: HttpRequest) -> HttpResponse:
    logout(request)
    return redirect("login")


def _dashboard_day_annotations():
    return {f"n_{sv}": Count("id", filter=Q(status=sv)) for sv in _ORDER_STATUS_CODES}


@login_required
def dashboard(request: HttpRequest) -> HttpResponse:
    if not is_admin_level(request.user):
        try:
            merchant = request.user.merchant_profile
        except Exception:
            merchant = None
        if merchant is not None:
            return redirect("merchants:detail", merchant_id=merchant.pk)
        try:
            delegate = request.user.delegate_profile
        except Exception:
            delegate = None
        if delegate is not None:
            return redirect("delegates:detail", delegate_id=delegate.pk)

    qs = orders_queryset_for_user(request.user)
    qs = _apply_order_filters(qs, request)
    base = qs
    ag = base.aggregate(
        total_orders=Count("id"),
        total_shipping=Sum("shipping_price"),
    )
    agg_by_status_rows = (
        base.values("status")
        .annotate(c=Count("id"), amt=Sum("product_price"))
    )
    count_by_status = {r["status"]: r["c"] or 0 for r in agg_by_status_rows}
    amt_by_status = {r["status"]: r["amt"] or 0 for r in agg_by_status_rows}

    d = {"c": count_by_status.get(OrderStatus.DELIVERED, 0), "amt": amt_by_status.get(OrderStatus.DELIVERED, 0)}
    t = {"c": count_by_status.get(OrderStatus.IN_TRANSIT, 0), "amt": amt_by_status.get(OrderStatus.IN_TRANSIT, 0)}
    a = {"c": count_by_status.get(OrderStatus.ACCOUNTED, 0), "amt": amt_by_status.get(OrderStatus.ACCOUNTED, 0)}

    status_cards = []
    for sv in _ORDER_STATUS_CODES:
        label = _ORDER_STATUS_LABELS.get(str(sv), str(sv))
        status_cards.append(
            {
                "code": sv,
                "label": label,
                "count": count_by_status.get(str(sv), 0),
                "product_amt": amt_by_status.get(str(sv), 0),
                "hex_color": _CHART_HEX.get(sv, "#6c757d"),
            }
        )

    secondary_status_cards = [row for row in status_cards if str(row["code"]) not in {str(k) for k in _DASH_LEGACY_THREE}]

    agg_dates = base.aggregate(mn=Min("order_date"), mx=Max("order_date"))
    ann_kw = _dashboard_day_annotations()
    chart_labels = []
    datasets = []
    chart_has_data = False
    if agg_dates["mx"] is None:
        by_day_list = []
    else:
        chart_end = agg_dates["mx"]
        mn = agg_dates["mn"] or chart_end
        window_start = max(mn, chart_end - timedelta(days=13))
        by_day_qs = base.filter(order_date__gte=window_start, order_date__lte=chart_end)
        by_day_list = list(by_day_qs.values("order_date").annotate(**ann_kw).order_by("order_date"))
        chart_labels = [str(row["order_date"]) for row in by_day_list]
        for sv in _ORDER_STATUS_CODES:
            key = f"n_{sv}"
            datasets.append(
                {
                    "label": _ORDER_STATUS_LABELS.get(str(sv), str(sv)),
                    "backgroundColor": _CHART_HEX.get(sv, "#6c757d"),
                    "data": [row.get(key) or 0 for row in by_day_list],
                }
            )
        chart_has_data = any((row.get(f"n_{sv}") or 0) > 0 for row in by_day_list for sv in _ORDER_STATUS_CODES)

    return render(
        request,
        "orders/dashboard.html",
        {
            "delivered_orders_count": d["c"] or 0,
            "in_transit_orders_count": t["c"] or 0,
            "accounted_orders_count": a["c"] or 0,
            "total_delivered_amount": d["amt"] or 0,
            "total_in_transit_amount": t["amt"] or 0,
            "total_accounted_amount": a["amt"] or 0,
            "total_orders": ag["total_orders"] or 0,
            "total_shipping": ag["total_shipping"] or 0,
            "secondary_status_cards": secondary_status_cards,
            "chart_labels_json": json.dumps(chart_labels, ensure_ascii=False),
            "chart_series_json": json.dumps(datasets, ensure_ascii=False),
            "chart_has_data": chart_has_data,
            "chart_ready_json": json.dumps(chart_has_data),
        },
    )


@login_required
def order_list(request: HttpRequest) -> HttpResponse:
    qs = orders_queryset_for_user(request.user)
    qs = _apply_order_filters(qs, request).select_related("merchant", "delegate", "branch")

    paginator = Paginator(qs.order_by("-order_date", "-id"), 25)
    page_obj = paginator.get_page(query_first_str(request.GET, "page") or 1)

    qcopy = request.GET.copy()
    qcopy.pop("page", None)
    filter_qs = qcopy.urlencode()

    return render(
        request,
        "orders/order_list.html",
        {
            "page_obj": page_obj,
            "filter_qs": filter_qs,
        },
    )


@login_required
def order_add(request: HttpRequest) -> HttpResponse:
    if request.method == "POST":
        form = OrderForm(request.POST)
        if form.is_valid():
            order = form.save(commit=False)
            order.created_by = request.user
            order.updated_by = request.user
            order.owned_by = request.user
            setattr(order, "_save_user", request.user)
            order.save()
            form.save_m2m()
            log_action(request, AuditAction.CREATE, order, "إضافة شحنة")
            messages.success(request, "تم حفظ الشحنة.")
            return redirect("orders:list")
    else:
        last = orders_queryset_for_user(request.user).order_by("-id").first()
        initial = {}
        if last:
            initial = {
                "merchant": last.merchant_id,
                "delegate": last.delegate_id,
                "shipping_price": last.shipping_price,
                "order_date": last.order_date,
                "branch": last.branch_id,
                "customer_name": last.customer_name or "",
                "customer_phone": last.customer_phone or "",
            }
            if last.product_price:
                initial["product_price"] = last.product_price
        form = OrderForm(initial=initial)
    return render(
        request,
        "orders/order_form.html",
        {
            "form": form,
            "title": "إضافة شحنة",
            "can_manage_staff_records": can_manage_staff_records(request.user),
            "merchant_edit_tpl": reverse("merchants:edit", kwargs={"merchant_id": STUB_PK}),
            "delegate_edit_tpl": reverse("delegates:edit", kwargs={"delegate_id": STUB_PK}),
        },
    )


@login_required
def order_edit(request: HttpRequest, pk: int) -> HttpResponse:
    order = get_object_or_404(orders_queryset_for_user(request.user), pk=pk)
    if request.method == "POST":
        form = OrderForm(request.POST, instance=order)
        if form.is_valid():
            o = form.save(commit=False)
            o.updated_by = request.user
            setattr(o, "_save_user", request.user)
            o.save()
            form.save_m2m()
            log_action(request, AuditAction.UPDATE, o, "تعديل شحنة")
            messages.success(request, "تم حفظ التعديلات.")
            return redirect("orders:list")
    else:
        form = OrderForm(instance=order)
    settlement_preview = compute_order_settlement(order)
    return render(
        request,
        "orders/order_form.html",
        {
            "form": form,
            "title": "تعديل شحنة",
            "edit_order": order,
            "settlement_preview": settlement_preview,
            "can_manage_staff_records": can_manage_staff_records(request.user),
            "merchant_edit_tpl": reverse("merchants:edit", kwargs={"merchant_id": STUB_PK}),
            "delegate_edit_tpl": reverse("delegates:edit", kwargs={"delegate_id": STUB_PK}),
        },
    )


@login_required
def order_soft_delete(request: HttpRequest, pk: int) -> HttpResponse:
    if not is_super_admin(request.user):
        messages.error(request, "حذف الشحنات متاح للمدير العام فقط لتفادي فقد البيانات.")
        return redirect("orders:list")
    order = get_object_or_404(orders_queryset_for_user(request.user), pk=pk)
    if request.method != "POST":
        return redirect("orders:list")
    try:
        create_sqlite_backup("pre_order_delete")
    except Exception:
        pass
    wb = order.waybill_number
    order.soft_delete(request.user)
    log_action(request, AuditAction.DELETE, order, f"حذف منطقي للشحنة {wb}")
    messages.success(request, "تم إخفاء الشحنة من القائمة (حذف منطقي).")
    return redirect("orders:list")


def _parse_bulk_ids(request: HttpRequest) -> list[int]:
    if request.method == "POST":
        raw_ids = request.POST.getlist("order_ids")
        if raw_ids:
            out = []
            for x in raw_ids:
                try:
                    out.append(int(x))
                except ValueError:
                    continue
            return out
    raw = (request.GET.get("ids") or "").strip()
    out = []
    for part in raw.split(","):
        part = part.strip()
        if not part:
            continue
        try:
            out.append(int(part))
        except ValueError:
            continue
    return out


def _bulk_queryset(request: HttpRequest, ids: list[int]):
    if not ids:
        return orders_queryset_for_user(request.user).none()
    return orders_queryset_for_user(request.user).filter(pk__in=ids)


@login_required
def order_bulk(request: HttpRequest) -> HttpResponse:
    if request.method != "POST":
        return redirect("orders:list")
    ids = _parse_bulk_ids(request)
    action = query_first_str(request.POST, "action")
    qs = _bulk_queryset(request, ids)
    if not ids:
        messages.error(request, "لم يتم اختيار أي شحنة.")
        return redirect("orders:list")

    if action == "excel":
        return redirect(f"{reverse('orders:export')}?ids={','.join(str(i) for i in ids)}")
    if action == "print":
        return redirect(f"{reverse('orders:print')}?ids={','.join(str(i) for i in ids)}")
    if action == "qr":
        return redirect(f"{reverse('orders:qr')}?ids={','.join(str(i) for i in ids)}")
    if action == "whatsapp":
        body = build_bulk_whatsapp_text_ar(qs.order_by("waybill_number").select_related("merchant", "delegate"))
        url = "https://wa.me/?text=" + quote(body)
        return HttpResponseRedirect(url)
    if action == "bulk_status" and is_admin_level(request.user):
        new_status = query_first_str(request.POST, "new_status")
        if new_status not in _ORDER_STATUS_CODES:
            messages.error(request, "حالة غير صالحة.")
            return redirect("orders:list")
        n = 0
        for o in qs:
            if o.status == new_status:
                continue
            o.status = new_status
            o.updated_by = request.user
            setattr(o, "_save_user", request.user)
            o.save(update_fields=["status", "updated_at", "updated_by"])
            n += 1
        log_action(request, AuditAction.UPDATE, None, f"تعديل جماعي لحالة {n} شحنة")
        messages.success(request, f"تم تحديث {n} شحنة.")
        return redirect("orders:list")

    messages.error(request, "إجراء غير معروف أو غير مسموح.")
    return redirect("orders:list")


@login_required
def order_export_excel(request: HttpRequest) -> HttpResponse:
    ids = _parse_bulk_ids(request)
    qs = _bulk_queryset(request, ids).select_related("merchant", "delegate", "branch").order_by("-order_date", "-id")
    if not ids:
        messages.error(request, "لا توجد شحنات للتصدير.")
        return redirect("orders:list")

    wb = Workbook()
    ws = wb.active
    if ws is None:
        ws = wb.create_sheet(title="orders")
    else:
        ws.title = "orders"
    headers = [
        "التاريخ",
        "البوليصة",
        "الفرع",
        "التاجر",
        "هاتف التاجر",
        "المندوب",
        "هاتف المندوب",
        "العميل",
        "هاتف العميل",
        "العنوان",
        "سعر المنتج",
        "سعر الشحن",
        "الحالة",
    ]
    ws.append(headers)
    for o in qs:
        ws.append(
            [
                o.order_date.isoformat(),
                o.waybill_number,
                o.branch.name if o.branch else "",
                o.merchant.name if o.merchant else "",
                o.merchant.phone if o.merchant else "",
                o.delegate.name if o.delegate else "",
                o.delegate.phone if o.delegate else "",
                o.customer_name,
                o.customer_phone,
                o.customer_address,
                float(o.product_price),
                float(o.shipping_price),
                _ORDER_STATUS_LABELS.get(o.status, str(o.status)),
            ]
        )
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 16

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    resp = HttpResponse(
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    resp["Content-Disposition"] = 'attachment; filename="orders_export.xlsx"'
    return resp


@login_required
def order_print(request: HttpRequest) -> HttpResponse:
    ids = _parse_bulk_ids(request)
    qs = _bulk_queryset(request, ids).select_related("merchant", "delegate").order_by("-order_date", "-id")
    return render(request, "orders/order_print.html", {"orders": qs})


@login_required
def order_qr(request: HttpRequest) -> HttpResponse:
    ids = _parse_bulk_ids(request)
    qs = _bulk_queryset(request, ids).select_related("merchant", "delegate")
    rows = []
    for o in qs:
        rows.append(
            {
                "order": o,
                "qr_png_url": reverse("orders:qr_png", args=[o.pk]),
            }
        )
    return render(request, "orders/order_qr.html", {"rows": rows})


@login_required
def order_qr_png(request: HttpRequest, pk: int) -> HttpResponse:
    order = get_object_or_404(
        orders_queryset_for_user(request.user).select_related("merchant", "delegate"),
        pk=pk,
    )
    payload = sign_order_pk(order.pk)
    img = qrcode.make(payload)
    buf = BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    resp = HttpResponse(buf.getvalue(), content_type="image/png")
    resp["Cache-Control"] = "no-store"
    return resp


@login_required
def order_scanner(request: HttpRequest) -> HttpResponse:
    return render(request, "orders/order_scanner.html")


@login_required
def order_lookup_api(request: HttpRequest) -> HttpResponse:
    q = query_first_str(request.GET, "q")
    if not q:
        return JsonResponse({"ok": False, "error": "empty"}, status=400)
    order = None
    signed_pk = unsign_order_token(q.strip())
    if signed_pk is not None:
        order = orders_queryset_for_user(request.user).filter(pk=signed_pk).first()
    token = q
    if order is None and "WB:" in q:
        for part in q.split("|"):
            if part.startswith("WB:"):
                token = part[3:]
                break
        order = orders_queryset_for_user(request.user).filter(
            Q(waybill_number__iexact=token) | Q(external_waybill__iexact=token)
        ).first()
    if order is None:
        order = orders_queryset_for_user(request.user).filter(
            Q(waybill_number__iexact=q.strip()) | Q(external_waybill__iexact=q.strip())
        ).first()
    if not order:
        return JsonResponse({"ok": False, "found": False})
    return JsonResponse(
        {
            "ok": True,
            "found": True,
            "id": order.pk,
            "waybill_number": order.waybill_number,
            "external_waybill": order.external_waybill,
            "customer_name": order.customer_name,
            "customer_phone": order.customer_phone,
            "customer_address": order.customer_address,
            "status": order.get_status_display(),
            "edit_url": reverse("orders:edit", args=[order.pk]),
        }
    )


@login_required
def order_import(request: HttpRequest) -> HttpResponse:
    media_root = Path(settings.MEDIA_ROOT)
    import_dir = media_root / "imports"
    import_dir.mkdir(parents=True, exist_ok=True)

    if request.method == "POST" and request.FILES.get("file"):
        form = OrderImportUploadForm(request.POST, request.FILES)

        if not form.is_valid():
            return render(
                request,
                "orders/order_import.html",
                {"step": "upload", "upload_form": form},
            )

        file = form.cleaned_data["file"]
        token = f"{uuid4()}.xlsx"
        path = import_dir / token

        with path.open("wb") as f:
            for chunk in file.chunks():
                f.write(chunk)

        wb = load_workbook(path, data_only=True)
        sheet = wb.active
        if sheet is None and wb.sheetnames:
            sheet = wb[wb.sheetnames[0]]
        if sheet is None:
            messages.error(request, "ملف Excel بدون أوراق عمل.")
            path.unlink(missing_ok=True)
            return render(
                request,
                "orders/order_import.html",
                {"step": "upload", "upload_form": OrderImportUploadForm()},
            )
        header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))

        mapping_form = OrderImportMappingForm(
            initial={"token": token},
            sheet_choices=[(s, s) for s in wb.sheetnames],
            column_choices=[(str(i), str(h)) for i, h in enumerate(header)],
        )

        return render(
            request,
            "orders/order_import.html",
            {
                "step": "map",
                "mapping_form": mapping_form,
                "token": token,
            },
        )

    if request.method == "POST" and query_first_str(request.POST, "token"):
        token = query_first_str(request.POST, "token")
        if not token or "/" in token or "\\" in token or ".." in token:
            messages.error(request, "رمز الملف غير صالح.")
            return redirect("orders:import")
        path = import_dir / token

        if not path.is_file():
            return render(
                request,
                "orders/order_import.html",
                {
                    "step": "upload",
                    "upload_form": OrderImportUploadForm(),
                },
            )

        wb = load_workbook(path, data_only=True)
        form = OrderImportMappingForm(request.POST)

        if not form.is_valid():
            return render(
                request,
                "orders/order_import.html",
                {
                    "step": "map",
                    "mapping_form": form,
                    "token": token,
                },
            )

        sheet = wb[form.cleaned_data["sheet_name"]]
        created = 0
        updated = 0

        for row in sheet.iter_rows(min_row=2, values_only=True):
            raw_waybill = _to_str(
                _get(row, int(form.cleaned_data["waybill_number_col"]))
            )
            waybill = raw_waybill if (raw_waybill and WAYBILL_RE.match(raw_waybill)) else ""
            external_waybill = "" if waybill else (raw_waybill or None)

            merchant_name = _to_str(
                _get(row, int(form.cleaned_data["merchant_name_col"]))
            )
            if not merchant_name:
                continue
            merchant_phone = _to_str(
                _get(row, int(form.cleaned_data["merchant_phone_col"]))
            )
            merchant_defaults = {"phone": merchant_phone} if merchant_phone else {}
            merchant, _ = Merchant.objects.get_or_create(
                name=merchant_name,
                defaults=merchant_defaults,
            )
            if merchant_phone and not merchant.phone:
                merchant.phone = merchant_phone
                merchant.save(update_fields=["phone"])

            delegate_name = _to_str(
                _get(row, int(form.cleaned_data["delegate_name_col"]))
            )
            if not delegate_name:
                continue
            delegate_phone = _to_str(
                _get(row, int(form.cleaned_data["delegate_phone_col"]))
            )
            delegate_defaults = {"phone": delegate_phone} if delegate_phone else {}
            delegate, _ = Delegate.objects.get_or_create(
                name=delegate_name,
                defaults=delegate_defaults,
            )
            if delegate_phone and not delegate.phone:
                delegate.phone = delegate_phone
                delegate.save(update_fields=["phone"])

            try:
                order_payload = {
                    "order_date": _to_date(
                        _get(row, int(form.cleaned_data["order_date_col"]))
                    ),
                    "merchant": merchant,
                    "delegate": delegate,
                    "customer_name": _to_str(
                        _get(row, int(form.cleaned_data["customer_name_col"]))
                    ),
                    "customer_phone": _to_str(
                        _get(row, int(form.cleaned_data["customer_phone_col"]))
                    ),
                    "customer_address": _to_str(
                        _get(row, int(form.cleaned_data["customer_address_col"]))
                    ),
                    "shipping_price": _to_decimal(
                        _get(row, int(form.cleaned_data["shipping_price_col"]))
                    ),
                    "product_price": _to_decimal(
                        _get(row, int(form.cleaned_data["product_price_col"]))
                    ),
                    "notes": _to_str(_get(row, int(form.cleaned_data["notes_col"]))),
                    "external_waybill": external_waybill,
                    "updated_by": request.user,
                }

                if waybill:
                    o, was_created = Order.objects.get_or_create(
                        waybill_number=waybill,
                        defaults={
                            **order_payload,
                            "created_by": request.user,
                            "owned_by": request.user,
                        },
                    )
                    if was_created:
                        created += 1
                    else:
                        for k, v in order_payload.items():
                            setattr(o, k, v)
                        setattr(o, "_save_user", request.user)
                        o.save()
                        updated += 1
                else:
                    o = Order(
                        waybill_number="",
                        created_by=request.user,
                        owned_by=request.user,
                        **order_payload,
                    )
                    setattr(o, "_save_user", request.user)
                    o.save()
                    created += 1
            except Exception:
                continue

        path.unlink(missing_ok=True)

        return render(
            request,
            "orders/order_import.html",
            {
                "step": "done",
                "created": created,
                "updated": updated,
            },
        )

    return render(
        request,
        "orders/order_import.html",
        {
            "step": "upload",
            "upload_form": OrderImportUploadForm(),
        },
    )
